"""Forward-operator metrics table as an Excel sheet (same content and colours as forward_metrics_table.png).
Values: erp_forward/plots/models/100k/ALL_MODELS/forward_models_metrics.csv (200-epoch checkpoints, 100k dataset, 10,000 test spectra).
Colours of the four metric columns are conditional formats (best value green; worse than halfway between best and worst red), so they follow the numbers if edited.
Output: forward_metrics_table.xlsx
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = Path(__file__).resolve().parent / "forward_metrics_table.xlsx"

# group colours (fill, border): blue = DeepONet family and NN, violet = transform operators, orange = graph / attention, yellow = sine
BLUE, VIOLET, ORANGE, YELLOW = ("DCE8F6", "8FB3DD"), ("E8DFF5", "B39DDB"), ("FFE3D1", "F0A678"), ("FBF2C4", "E6C84A")
ROWS = [  # architecture, core operation, RMSE, R2, corr, peak RMSE, group
    ("DeepONet (DON)", "Inner product", 6.43, 0.589, 0.842, 12.92, BLUE),
    ("Deep Neural Operator (DNO)", "Modulation (FiLM)", 2.92, 0.915, 0.963, 6.78, BLUE),
    ("Deep Cat Operator (DCO)", "Concatenation", 2.70, 0.927, 0.967, 6.70, BLUE),
    ("Fourier Neural Operator (FNO)", "Fourier transform", 3.45, 0.882, 0.945, 7.96, VIOLET),
    ("Wavelet Neural Operator (WNO)", "Wavelet transform", 3.46, 0.881, 0.943, 9.05, VIOLET),
    ("Laplace Neural Operator (LNO)", "Poles and residues", 3.18, 0.899, 0.957, 7.15, VIOLET),
    ("Graph Neural Operator (GNO)", "Graph network", 3.84, 0.853, 0.949, 6.57, ORANGE),
    ("Set Transformer Operator (STO)", "Attention", 3.86, 0.852, 0.947, 7.03, ORANGE),
    ("SIREN Neural Operator (SIREN)", "Sine layers", 3.12, 0.903, 0.958, 6.96, YELLOW),
    ("Plain Neural Network (NN)", "Dense layers", 8.53, 0.275, 0.764, 15.66, BLUE),
]
HEAD = ["Architecture", "Core operation", "RMSE (dB)", "R²", "Correlation", "Peak RMSE (dB)"]
FONT = "Arial"

wb = Workbook()
ws = wb.active
ws.title = "Forward operators"
thin = lambda c: Side(style="thin", color=c)
grey = thin("C9CED6")

for j, h in enumerate(HEAD, 1):
    c = ws.cell(row=1, column=j, value=h)
    c.font = Font(name=FONT, bold=True, size=12)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[1].height = 34

for i, (arch, core, rmse, r2, corr, peak, (fill, edge)) in enumerate(ROWS, 2):
    for j, v in enumerate((arch, core), 1):
        c = ws.cell(row=i, column=j, value=v)
        c.fill = PatternFill("solid", fgColor=fill)
        c.border = Border(*(thin(edge),) * 4)
        c.font = Font(name=FONT, size=11, color="1F2933")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for j, (v, fmt) in enumerate(((rmse, "0.00"), (r2, "0.000"), (corr, "0.000"), (peak, "0.00")), 3):
        c = ws.cell(row=i, column=j, value=v)
        c.number_format = fmt
        c.border = Border(*(grey,) * 4)
        c.font = Font(name=FONT, size=11, color="1F2933")
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[i].height = 22

first, last = 2, 1 + len(ROWS)
green = PatternFill("solid", bgColor="D9EFD9", fgColor="D9EFD9")
red = PatternFill("solid", bgColor="F8C8C8", fgColor="F8C8C8")
for col, lower_is_better in (("C", True), ("D", False), ("E", False), ("F", True)):
    rng = f"{col}{first}:{col}{last}"
    col_range = f"{col}${first}:{col}${last}"
    best, worst = ("MIN", "MAX") if lower_is_better else ("MAX", "MIN")
    cmp = ">" if lower_is_better else "<"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"{col}{first}={best}({col_range})"], fill=green, stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"{col}{first}{cmp}(MIN({col_range})+MAX({col_range}))/2"], fill=red))

note = last + 2
ws.cell(row=note, column=1, value="200-epoch models, 100k dataset (10,000 test spectra). Correlation: Pearson r over all points. Peak RMSE: error at the resonance peaks of the true ERP.")
ws.cell(row=note + 1, column=1, value="Colours: green = best value in the column; red = worse than halfway between the best and the worst value. Source: erp_forward/plots/models/100k/ALL_MODELS/forward_models_metrics.csv (values typed in, 2-3 decimals).")
for r in (note, note + 1):
    ws.cell(row=r, column=1).font = Font(name=FONT, size=10, italic=True, color="52606D")

for col, w in zip("ABCDEF", (34, 22, 13, 11, 14, 16)):
    ws.column_dimensions[col].width = w
ws.sheet_view.showGridLines = False
wb.save(OUT)
print("saved", OUT)
